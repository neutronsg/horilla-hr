import base64
import io
import tempfile
import uuid
import zipfile
from datetime import date, datetime, time
from types import SimpleNamespace
from unittest.mock import MagicMock, patch
from zoneinfo import ZoneInfo

from django.contrib.auth.models import Group, Permission
from django.contrib.messages.storage.fallback import FallbackStorage
from django.contrib.sessions.backends.db import SessionStore
from django.core.files.base import ContentFile
from django.core.files.storage import Storage
from django.http import HttpResponse
from django.test import RequestFactory, TestCase, override_settings

from base.models import CompanyGroupAssignment, DefaultExportPermission
from horilla.horilla_middlewares import _thread_locals
from horilla.testkit import (
    clear_selected_company,
    make_company,
    make_employee,
    make_user,
)
from payroll.claim_reminders import (
    approval_recipients,
    run_due_claim_reminders,
    submission_cutoff,
)
from payroll.hr_access import can_manage_claim_reminders
from payroll.hr_delivery_models import (
    ClaimReminderDelivery,
    ClaimReminderPolicy,
    PayslipEmailDelivery,
)
from payroll.hr_delivery_views import (
    ClaimReminderForm,
    claim_reminder_settings,
    payslip_email_results,
)
from payroll.models.models import (
    Payslip,
    Reimbursement,
    ReimbursementMultipleAttachment,
)
from payroll.threadings.mail import MailSendThread
from report.claim_exports import claim_queryset, export_claim_bundle
from report.engine import ReportFilters
from report.metrics.payroll import reimbursement_register


@override_settings(COMPANY_SCOPED_PERMISSIONS=True, TIME_ZONE="Asia/Singapore")
class HRDeliveryTests(TestCase):
    def setUp(self):
        clear_selected_company()
        self.company = make_company("Delivery QA", country="Singapore")
        self.other = make_company("Other QA")
        self.hr = make_employee(
            company=self.company, email="zeyu+qa-hr@neutron.sg", first_name="HR"
        )
        self.employee = make_employee(
            company=self.company, email="zeyu+qa-employee@neutron.sg"
        )
        self.foreign = make_employee(
            company=self.other, email="zeyu+qa-other@neutron.sg"
        )
        self.grant(self.hr, "HR Manager", self.company)
        self.request = self.req(self.hr.employee_user_id)
        self.policy = ClaimReminderPolicy.objects.create(company=self.company)
        self.now = datetime(2026, 10, 22, 9, 10, tzinfo=ZoneInfo("Asia/Singapore"))
        self.temp = tempfile.TemporaryDirectory()
        self.storage_settings = override_settings(MEDIA_ROOT=self.temp.name)
        self.storage_settings.enable()

    def tearDown(self):
        _thread_locals.request = None
        clear_selected_company()
        self.storage_settings.disable()
        self.temp.cleanup()
        super().tearDown()

    def grant(
        self,
        employee,
        name,
        company,
        permissions=(
            "view_employee",
            "change_employee",
            "view_reimbursement",
            "change_reimbursement",
            "view_payslip",
            "add_payslip",
        ),
    ):
        group, _ = Group.objects.get_or_create(name=name)
        group.permissions.set(Permission.objects.filter(codename__in=permissions))
        CompanyGroupAssignment.objects.get_or_create(
            user=employee.employee_user_id, company=company, group=group
        )
        employee.employee_user_id.groups.add(group)
        return group

    def req(self, user, path="/", data=None, method="get", **headers):
        request = getattr(RequestFactory(), method)(path, data or {}, **headers)
        request.user = user
        request.session = SessionStore()
        request.session["selected_company"] = "all"
        request.session.save()
        request._messages = FallbackStorage(request)
        _thread_locals.request = request
        return request

    def claim(self, employee=None, **kwargs):
        claim = Reimbursement(
            employee_id=employee or self.employee,
            title="Taxi",
            description="Client meeting",
            allowance_on=date(2026, 9, 30),
            amount=10,
            **kwargs,
        )
        # Avoid reimbursement save side effects: these fixtures test export/notification snapshots.
        Reimbursement.objects.bulk_create([claim])
        return claim

    def slip(self, employee=None, status="confirmed"):
        self.slip_count = getattr(self, "slip_count", 0) + 1
        month = 8 + self.slip_count
        return Payslip.objects.create(
            employee_id=employee or self.employee,
            status=status,
            start_date=date(2026, month, 1),
            end_date=date(2026, month, 28),
            gross_pay=1000,
            net_pay=1000,
            deduction=0,
            pay_head_data={},
        )

    def filters(self, **kwargs):
        return ReportFilters(
            date(2026, 10, 1),
            date(2026, 10, 31),
            company_id=self.company.pk,
            employment_status="all",
            request=self.request,
            **kwargs,
        )

    def test_hr_role_and_permission_must_match_target_company(self):
        self.assertTrue(
            can_manage_claim_reminders(self.hr.employee_user_id, self.company.pk)
        )
        self.assertFalse(
            can_manage_claim_reminders(self.hr.employee_user_id, self.other.pk)
        )
        self.assertFalse(
            can_manage_claim_reminders(self.employee.employee_user_id, self.company.pk)
        )
        self.grant(self.foreign, "Manager", self.other)
        self.assertFalse(
            can_manage_claim_reminders(self.foreign.employee_user_id, self.other.pk)
        )

    def test_role_recipients_are_company_scoped_active_and_approval_capable(self):
        self.grant(self.foreign, "HR Manager", self.other)
        admin = make_employee(company=self.company, email="zeyu+qa-admin@neutron.sg")
        self.grant(admin, "Admin", self.company)
        self.assertEqual([e.pk for e in approval_recipients(self.policy)], [self.hr.pk])
        self.policy.include_admin = True
        self.assertCountEqual(
            [e.pk for e in approval_recipients(self.policy)], [self.hr.pk, admin.pk]
        )
        self.hr.employee_user_id.is_active = False
        self.hr.employee_user_id.save(update_fields=["is_active"])
        self.assertEqual([e.pk for e in approval_recipients(self.policy)], [admin.pk])

    def test_settings_existing_hr_access_and_no_cross_company_write(self):
        request = self.req(
            self.hr.employee_user_id,
            "/payroll/claim-reminders/",
            {"company_id": self.company.pk},
        )
        self.assertEqual(claim_reminder_settings(request).status_code, 200)
        request = self.req(
            self.hr.employee_user_id,
            "/payroll/claim-reminders/",
            {"company_id": self.other.pk},
            method="post",
        )
        self.assertEqual(claim_reminder_settings(request).status_code, 403)
        self.assertFalse(
            ClaimReminderPolicy.objects.filter(company=self.other).exists()
        )

    def test_invalid_cutoff_and_lead_time_rejected(self):
        form = ClaimReminderForm(
            {
                "cutoff_day": 32,
                "days_before": 15,
                "submission_time": "09:00",
                "approval_time": "09:00",
                "timezone": "Asia/Singapore",
            }
        )
        self.assertFalse(form.is_valid())
        self.assertIn("cutoff_day", form.errors)
        self.assertIn("days_before", form.errors)

    def test_cutoff_handles_february_and_reminder_previous_month(self):
        self.policy.cutoff_day = 31
        self.policy.days_before = 3
        self.assertEqual(
            submission_cutoff(self.policy, datetime(2026, 2, 25)), date(2026, 2, 28)
        )
        self.policy.cutoff_day = 2
        self.assertEqual(
            submission_cutoff(self.policy, datetime(2026, 12, 30)), date(2027, 1, 2)
        )

    @patch("payroll.claim_reminders.EmailMessage.send", return_value=1)
    def test_disabled_reminders_do_not_send(self, send):
        self.claim()
        self.assertEqual(run_due_claim_reminders(self.now), [])
        send.assert_not_called()

    @patch("payroll.claim_reminders.EmailMessage.send", return_value=1)
    def test_monthly_submission_deduplication_and_existing_claim_not_suppressed(
        self, send
    ):
        self.claim()
        self.policy.submission_enabled = True
        self.policy.save()
        results = run_due_claim_reminders(self.now)
        self.assertEqual(len(results), 2)
        self.assertTrue(all(d.status == "sent" for d in results))
        self.assertEqual(run_due_claim_reminders(self.now), [])
        self.assertEqual(send.call_count, 2)
        self.assertFalse(
            ClaimReminderDelivery.objects.filter(employee=self.foreign).exists()
        )

    @patch("payroll.claim_reminders.EmailMessage.send", return_value=1)
    def test_before_send_time_and_missed_date_do_not_send(self, send):
        self.policy.submission_enabled = True
        self.policy.save()
        run_due_claim_reminders(self.now.replace(hour=8))
        run_due_claim_reminders(self.now.replace(day=23))
        send.assert_not_called()

    @patch("payroll.claim_reminders.EmailMessage.send", return_value=0)
    def test_failed_submission_attempt_is_visible_and_not_duplicated(self, send):
        self.policy.submission_enabled = True
        self.policy.save()
        self.assertEqual(
            {d.status for d in run_due_claim_reminders(self.now)}, {"failed"}
        )
        run_due_claim_reminders(self.now)
        self.assertEqual(send.call_count, 2)

    @patch("payroll.claim_reminders.EmailMessage.send", return_value=1)
    def test_pending_digest_only_requested_company_claims_and_stops_after_approval(
        self, send
    ):
        claim = self.claim()
        self.claim(employee=self.foreign)
        self.policy.approval_enabled = True
        self.policy.save()
        results = run_due_claim_reminders(self.now)
        self.assertEqual(len(results), 1)
        self.assertEqual(results[0].employee_id, self.hr.pk)
        run_due_claim_reminders(self.now)
        self.assertEqual(send.call_count, 1)
        Reimbursement.objects.filter(pk=claim.pk).update(status="approved")
        self.assertEqual(run_due_claim_reminders(self.now.replace(day=23)), [])
        self.assertEqual(send.call_count, 1)

    def test_export_end_date_includes_evening_and_excludes_next_day(self):
        today = self.claim()
        nextday = self.claim()
        Reimbursement.objects.filter(pk=today.pk).update(
            created_at=datetime(2026, 10, 31, 23, 59, tzinfo=ZoneInfo("Asia/Singapore"))
        )
        Reimbursement.objects.filter(pk=nextday.pk).update(
            created_at=datetime(2026, 11, 1, 0, 0, tzinfo=ZoneInfo("Asia/Singapore"))
        )
        self.assertEqual(
            list(claim_queryset(self.filters()).values_list("pk", flat=True)),
            [today.pk],
        )

    def test_undated_claim_stays_visible_in_submitted_date_register(self):
        claim = self.claim()
        Reimbursement.objects.filter(pk=claim.pk).update(created_at=None)
        rows = reimbursement_register(self.filters())["table"]["rows"]
        self.assertEqual([row["claim_id"] for row in rows], [claim.pk])
        self.assertEqual(rows[0]["raised"], "")

    def test_payroll_date_basis_selects_september_claim_submitted_in_october(self):
        claim = self.claim()
        Reimbursement.objects.filter(pk=claim.pk).update(
            created_at=datetime(2026, 10, 2, tzinfo=ZoneInfo("Asia/Singapore"))
        )
        filters = self.filters(claim_date_basis="payroll")
        filters.from_date, filters.to_date = date(2026, 9, 1), date(2026, 9, 30)
        self.assertEqual(claim_queryset(filters).count(), 1)
        filters.claim_date_basis = "submitted"
        self.assertEqual(claim_queryset(filters).count(), 0)

    def test_totals_and_exports_cover_more_than_300_and_500_rows(self):
        Reimbursement.objects.bulk_create(
            [
                Reimbursement(
                    employee_id=self.employee,
                    title=f"Taxi {n}",
                    allowance_on=date(2026, 10, 1),
                    amount=10,
                )
                for n in range(510)
            ]
        )
        filters = self.filters()
        payload = reimbursement_register(filters)
        self.assertEqual(payload["kpis"][1]["value"], 510)
        self.assertEqual(payload["kpis"][2]["value"], 5100)
        self.assertEqual(len(payload["table"]["rows"]), 500)
        filters.extra["row_limit"] = None
        self.assertEqual(len(reimbursement_register(filters)["table"]["rows"]), 510)

    def test_claim_scope_applies_even_without_thread_company_context(self):
        self.claim()
        self.claim(employee=self.foreign)
        filters = self.filters()
        filters.company_id = None
        self.assertEqual(claim_queryset(filters).count(), 1)

    def test_claim_export_respects_company_export_policy_in_all_companies_session(self):
        from report.claim_exports import can_export_claims
        import report.metrics
        from report.registry import get_report

        self.claim()
        DefaultExportPermission.objects.create(
            company_id=self.company, is_enabled=False
        )
        self.assertFalse(
            can_export_claims(
                self.request, get_report("reimbursement-register"), self.filters()
            )
        )
        self.hr.employee_user_id.user_permissions.add(
            Permission.objects.get(
                content_type__app_label="payroll", codename="export_reimbursement"
            )
        )
        self.assertTrue(
            can_export_claims(
                self.request, get_report("reimbursement-register"), self.filters()
            )
        )

    def test_report_export_endpoint_returns_complete_xlsx_and_zip_with_matching_rows(
        self,
    ):
        from report.views.standard_reports import standard_report_export
        from openpyxl import load_workbook

        claim = self.claim()
        params = {
            "period_preset": "custom",
            "from_date": "2026-09-01",
            "to_date": "2026-09-30",
            "claim_date_basis": "payroll",
            "company_id": self.company.pk,
        }
        request = self.req(
            self.hr.employee_user_id,
            "/report/standard/reimbursement-register/export/",
            {**params, "format": "xlsx"},
        )
        response = standard_report_export(request, "reimbursement-register")
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(response.content))
        self.assertIn("Data", workbook.sheetnames)
        request = self.req(
            self.hr.employee_user_id,
            "/report/standard/reimbursement-register/export/",
            {**params, "format": "zip"},
        )
        response = standard_report_export(request, "reimbursement-register")
        self.assertEqual(response["Content-Type"], "application/zip")
        archive = zipfile.ZipFile(io.BytesIO(response.content))
        self.assertIn(str(claim.pk).encode(), archive.read("attachment-manifest.csv"))

    def test_private_receipt_storage_needs_neither_local_path_nor_public_url(self):
        claim = self.claim()
        claim.attachment.name = "private/receipt.png"
        claim.attachment.storage = NoPathStorage()
        response = export_claim_bundle(
            {"table": {"rows": [], "columns": []}}, [claim], {}
        )
        self.assertEqual(response["X-Export-Partial"], "0")
        archive = zipfile.ZipFile(io.BytesIO(response.content))
        receipt = next(n for n in archive.namelist() if n.startswith("receipts/"))
        self.assertTrue(archive.read(receipt).startswith(b"\x89PNG"))

    def test_stopping_policy_prevents_a_reserved_poll_from_sending(self):
        from payroll.claim_reminders import deliver_once

        self.policy.submission_enabled = True
        self.policy.save()
        ClaimReminderPolicy.objects.filter(pk=self.policy.pk).update(
            submission_enabled=False
        )
        with patch("payroll.claim_reminders.EmailMessage.send") as send:
            self.assertIsNone(
                deliver_once(
                    self.policy,
                    "submission",
                    "2026-10",
                    self.employee,
                    "Reminder",
                    "Message",
                )
            )
            send.assert_not_called()

    def test_model_permissions_do_not_add_a_separate_hr_subsection_permission(self):
        self.assertFalse(
            Permission.objects.filter(
                content_type__app_label="payroll",
                content_type__model__in=[
                    "claimreminderpolicy",
                    "claimreminderdelivery",
                    "payslipemaildelivery",
                ],
            ).exists()
        )

    def test_zip_original_files_deduplicated_within_claim_and_manifest_missing_files(
        self,
    ):
        claim = self.claim()
        claim.attachment.save(
            "receipt.pdf", ContentFile(b"%PDF-test-receipt"), save=False
        )
        Reimbursement.objects.filter(pk=claim.pk).update(
            attachment=claim.attachment.name
        )
        duplicate = ReimbursementMultipleAttachment.objects.create(
            attachment=ContentFile(b"%PDF-test-receipt", name="same.pdf")
        )
        missing = ReimbursementMultipleAttachment.objects.create(
            attachment="not-found.pdf"
        )
        different = ReimbursementMultipleAttachment.objects.create(
            attachment=ContentFile(b"another-receipt", name="receipt.pdf")
        )
        claim.other_attachments.add(duplicate, missing, different)
        payload = reimbursement_register(self.filters())
        response = export_claim_bundle(payload, [claim], {})
        self.assertEqual(response["X-Export-Partial"], "1")
        self.assertIn("PARTIAL", response["Content-Disposition"])
        archive = zipfile.ZipFile(io.BytesIO(response.content))
        files = [n for n in archive.namelist() if n.startswith("receipts/")]
        self.assertEqual(len(files), 2)
        self.assertTrue(all(f"claim-{claim.pk}/" in n and ".." not in n for n in files))
        self.assertIn(b"duplicate omitted", archive.read("attachment-manifest.csv"))
        self.assertIn(b"MISSING", archive.read("attachment-manifest.csv"))
        self.assertTrue(archive.read("claims.xlsx").startswith(b"PK"))

    def test_payslip_partial_failure_updates_only_successful_employee(self):
        first = self.slip()
        second = self.slip(employee=self.hr)
        third = self.slip(employee=self.foreign)
        batch = uuid.uuid4()
        records = {}
        for slip in (first, second, third):
            delivery = PayslipEmailDelivery.objects.create(
                batch=batch,
                actor=self.request.user,
                employee=slip.employee_id,
                payslip_ids=[slip.pk],
            )
            records[slip.pk] = {
                "employee_id": slip.employee_id,
                "instances": [slip],
                "delivery": delivery,
            }
        with patch(
            "payroll.threadings.mail.payslip_pdf",
            return_value=HttpResponse(b"%PDF-test", content_type="application/pdf"),
        ), patch("payroll.threadings.mail.EmailMessage.send", side_effect=[1, 0]):
            MailSendThread(
                self.request, records, [first.pk, second.pk, third.pk]
            ).send_records()
        first.refresh_from_db()
        second.refresh_from_db()
        third.refresh_from_db()
        self.assertTrue(first.sent_to_employee)
        self.assertFalse(second.sent_to_employee)
        self.assertFalse(third.sent_to_employee)
        self.assertEqual(
            list(
                PayslipEmailDelivery.objects.filter(batch=batch).values_list(
                    "status", flat=True
                )
            ),
            ["sent", "failed", "failed"],
        )

    def test_draft_and_failed_pdf_never_mark_sent(self):
        for status in ("draft", "confirmed"):
            slip = self.slip(status=status)
            with patch(
                "payroll.threadings.mail.payslip_pdf",
                return_value=HttpResponse("failed", status=500),
            ), patch("payroll.threadings.mail.EmailMessage.send") as send:
                MailSendThread(
                    self.request,
                    {1: {"employee_id": slip.employee_id, "instances": [slip]}},
                    [slip.pk],
                ).send_records()
                send.assert_not_called()
            slip.refresh_from_db()
            self.assertFalse(slip.sent_to_employee)

    def test_payslip_uses_verified_smtp_sender_and_hr_reply_address(self):
        slip = self.slip()
        from django.core.mail import EmailMessage

        backend = MagicMock()
        backend.dynamic_from_email_with_display_name = "Neutron HR <hr@neutron.sg>"
        backend.send_messages.return_value = 1
        with patch(
            "payroll.threadings.mail.ConfiguredEmailBackend", return_value=backend
        ), patch(
            "payroll.threadings.mail.payslip_pdf",
            return_value=HttpResponse(b"%PDF-test"),
        ):
            MailSendThread(
                self.request,
                {1: {"employee_id": slip.employee_id, "instances": [slip]}},
                [slip.pk],
            ).send_records()
        message = backend.send_messages.call_args.args[0][0]
        self.assertEqual(message.from_email, "Neutron HR <hr@neutron.sg>")
        self.assertEqual(message.reply_to, [self.hr.get_mail()])

    def test_send_endpoint_skips_draft_and_result_page_rechecks_access(self):
        from payroll.views.component_views import send_slip

        draft = self.slip(status="draft")
        request = self.req(
            self.hr.employee_user_id,
            "/payroll/send-slip/",
            {"id": draft.pk},
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )
        with patch("payroll.views.component_views.MailSendThread.start") as start:
            response = send_slip(request)
            self.assertEqual(response.status_code, 200)
            start.assert_not_called()
        delivery = PayslipEmailDelivery.objects.get()
        self.assertEqual(delivery.status, "skipped")
        self.assertEqual(
            payslip_email_results(request, delivery.batch).status_code, 200
        )
        request = self.req(self.employee.employee_user_id)
        self.assertEqual(
            payslip_email_results(request, delivery.batch).status_code, 403
        )

    def test_result_page_only_failed_ids_in_retry_link(self):
        first, second = self.slip(), self.slip()
        batch = uuid.uuid4()
        PayslipEmailDelivery.objects.create(
            batch=batch,
            actor=self.request.user,
            employee=self.employee,
            payslip_ids=[first.pk],
            status="sent",
        )
        PayslipEmailDelivery.objects.create(
            batch=batch,
            actor=self.request.user,
            employee=self.employee,
            payslip_ids=[second.pk],
            status="failed",
        )
        response = payslip_email_results(self.request, batch)
        self.assertIn(f"?id={second.pk}".encode(), response.content)
        self.assertNotIn(f"?id={first.pk}".encode(), response.content)


class NoPathStorage(Storage):
    def _open(self, name, mode="rb"):
        return ContentFile(
            base64.b64decode(
                "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jF9kAAAAASUVORK5CYII="
            )
        )


class LeaveStorageEmailTests(TestCase):
    def setUp(self):
        clear_selected_company()
        self.company = make_company("Leave mail QA")
        employee = make_employee(
            company=self.company, email="zeyu+qa-leave-actor@neutron.sg"
        )
        self.request = RequestFactory().get("/")
        self.request.user = employee.employee_user_id
        self.request.session = {"selected_company": str(self.company.pk)}
        _thread_locals.request = self.request

    def tearDown(self):
        _thread_locals.request = None
        clear_selected_company()
        super().tearDown()

    def test_leave_email_does_not_require_private_logo_path(self):
        from leave.threading import LeaveMailSendThread

        self.company.icon.name = "private/logo.png"
        self.company.icon.storage = NoPathStorage()
        recipient = SimpleNamespace(
            get_company=lambda: self.company,
            get_mail=lambda: "zeyu+qa-leave@neutron.sg",
        )
        leave = SimpleNamespace(id=1, employee_id=self.request.user.employee_get)
        with patch(
            "leave.threading.EmailMultiAlternatives.send", return_value=1
        ) as send, patch(
            "leave.threading.render_to_string", return_value="<html>Leave</html>"
        ):
            LeaveMailSendThread(self.request, leave, "request").send_email(
                "Leave request", "Review", [recipient], 1
            )
            send.assert_called_once()
